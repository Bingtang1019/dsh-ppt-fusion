# Reopen a pptx with python-pptx and report what an independent reader sees.
#
# This is the structural half of the compatibility matrix (plan 3.14 / S16): it
# proves a package is readable by a second implementation without PowerPoint,
# which is what catches the relationship and content-type mistakes that WPS and
# older Office are least tolerant of. LibreOffice headless covers the rendering
# half.
#
# Usage: <venv python> scripts/probe-pptx-reopen.py <deck.pptx> [<deck.pptx> ...]
#
# Prints one JSON object per deck:
#   {path, ok, slides, textChars, shapes, tables, pictures, charts, error}
import json
import sys

from pptx import Presentation


def walk(shapes, record):
    """Count one shape collection, descending into groups.

    ppt-master emits group-heavy slides: a non-recursive reader sees 14 shapes
    and zero text on its native chart page, which looks like data loss but is
    only the reader stopping at the first level. The compatibility matrix needs
    the recursive numbers.
    """
    for shape in shapes:
        record["shapes"] += 1
        if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
            walk(shape.shapes, record)
            continue
        if shape.has_text_frame:
            record["textChars"] += len(shape.text_frame.text)
        if getattr(shape, "has_table", False):
            record["tables"] += 1
        if getattr(shape, "has_chart", False):
            record["charts"] += 1
        if shape.shape_type == 13 or getattr(shape, "image", None) is not None:
            record["pictures"] += 1


def describe(path):
    record = {"path": path, "ok": False, "slides": 0, "textChars": 0, "shapes": 0, "tables": 0, "pictures": 0, "charts": 0, "error": None}
    try:
        presentation = Presentation(path)
        record["slides"] = len(presentation.slides)
        for slide in presentation.slides:
            walk(slide.shapes, record)
        record["ok"] = record["slides"] > 0
    except Exception as exc:  # reported as data; this script never raises
        record["error"] = f"{type(exc).__name__}: {exc}"
    return record


for deck in sys.argv[1:]:
    print(json.dumps(describe(deck)))

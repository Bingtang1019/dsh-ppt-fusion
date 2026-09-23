#!/usr/bin/env python3
"""Extract a numeric design profile from one .pptx (V7.2 B1, ADR-061).

The output carries colours, fonts, sizes and geometry only: never a run's text,
a media file, or the input path. `--copy-media` is the one explicit opt-in that
copies the package's media into a caller-chosen directory, and it stays off by
default so a reference deck cannot follow the profile into a repository.

Run with the engine venv's Python (python-pptx), for example:

    python design-profile.py --input deck.pptx --output design-profile.json
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from pptx import Presentation

ROLES = ("cover", "toc", "section", "content", "ending")
EMU_PER_INCH = 914400
PROFILE_VERSION = 1


def to_hex(color) -> str | None:
    """@returns the run colour as uppercase #RRGGBB, or None when it is inherited."""
    try:
        if color is not None and color.type is not None:
            return f"#{str(color.rgb).upper()}"
    except (AttributeError, ValueError):
        return None
    return None


def luminance(hex_color: str) -> float:
    """@returns relative luminance in 0..1, enough to tell dark text from white/watermark."""
    value = hex_color.lstrip("#")[:6]
    if len(value) != 6:
        return 0.0
    channels = [int(value[index : index + 2], 16) / 255 for index in (0, 2, 4)]
    return round(0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2], 4)


def is_grey(hex_color: str) -> bool:
    """@returns True when the colour's channels differ by at most 0x08."""
    value = hex_color.lstrip("#")[:6]
    if len(value) != 6:
        return False
    channels = [int(value[index : index + 2], 16) for index in (0, 2, 4)]
    return max(channels) - min(channels) <= 8


def size_pt(run) -> float | None:
    """@returns the run's explicit font size in points, or None when it inherits."""
    size = run.font.size
    if size is None:
        return None
    return round(size.pt, 1)


def font_name(run) -> str | None:
    """@returns the run's explicit Latin/EA font name, or None when it inherits."""
    for attribute in ("name", "typeface"):
        try:
            value = getattr(run.font, attribute, None)
            if isinstance(value, str) and value != "":
                return value
        except (AttributeError, ValueError):
            continue
    try:
        if run.font._rPr is not None:  # noqa: SLF001 - python-pptx exposes no public EA accessor
            ea = run.font._rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")  # noqa: SLF001
            if ea is not None and ea.get("typeface"):
                return ea.get("typeface")
    except (AttributeError, ValueError):
        pass
    return None


def slide_runs(slide) -> list[dict]:
    """@returns every authored run on the slide with the geometry and style facts the profile needs."""
    runs: list[dict] = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text.strip() == "":
                    continue
                runs.append(
                    {
                        "text": run.text,
                        "sizePt": size_pt(run),
                        "bold": bool(run.font.bold),
                        "color": to_hex(run.font.color),
                        "font": font_name(run),
                        "x": shape.left,
                        "y": shape.top,
                        "w": shape.width,
                        "h": shape.height,
                    }
                )
    return runs


def runs_with_style(runs: list[dict]) -> list[dict]:
    """@returns only the runs that name a size, colour and font explicitly."""
    return [run for run in runs if run["sizePt"] is not None and run["color"] is not None and run["font"] is not None]


def has_full_bleed_picture(slide, width: int, height: int) -> bool:
    """@returns True when a picture covers at least 90% of the canvas."""
    for shape in slide.shapes:
        if shape.shape_type != 13:  # PICTURE
            continue
        if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
            continue
        if shape.width >= width * 0.9 and shape.height >= height * 0.9:
            return True
    return False


def classify_by_index(index: int, total: int) -> str:
    """@returns the index-based fallback role for a deck without the reference layout."""
    if index == 1:
        return "cover"
    if index == total:
        return "ending"
    return "content"


def is_number_like(text: str) -> bool:
    """@returns True for a numeral run such as `02`, `02.` or `12、`."""
    stripped = text.strip().strip(".．、 ")
    return stripped.isdigit() and len(stripped) <= 4


def map_roles(runs_per_slide: list[list[dict]], total: int, overrides: dict[int, str]) -> list[str]:
    """@returns one role per slide, preferring explicit overrides, then measured hints."""
    roles: list[str] = []
    for offset, runs in enumerate(runs_per_slide):
        index = offset + 1
        if index in overrides:
            roles.append(overrides[index])
            continue
        if index == 1:
            roles.append("cover")
            continue
        if index == total:
            roles.append("ending")
            continue
        watermark = any(run["sizePt"] is not None and run["sizePt"] >= 60 and run["color"] is not None and luminance(run["color"]) > 0.85 for run in runs)
        if watermark:
            roles.append("section")
            continue
        digits = [run for run in runs if run["sizePt"] is not None and run["sizePt"] >= 32 and is_number_like(run["text"])]
        if index == 2 and len(digits) >= 2:
            roles.append("toc")
            continue
        roles.append(classify_by_index(index, total))
    return roles


def parse_roles(spec: str | None) -> dict[int, str]:
    """@param spec - `role=1,2;role=3` or a JSON object mapping index to role."""
    if spec is None or spec.strip() == "":
        return {}
    text = spec.strip()
    overrides: dict[int, str] = {}
    if text.startswith("{"):
        for key, value in json.loads(text).items():
            if value not in ROLES:
                raise SystemExit(f"unknown role {value!r}")
            overrides[int(key)] = value
        return overrides
    for chunk in text.split(";"):
        if "=" not in chunk:
            raise SystemExit(f"bad --roles chunk {chunk!r}; use role=1,2;role=3")
        role, indices = chunk.split("=", 1)
        if role not in ROLES:
            raise SystemExit(f"unknown role {role!r}")
        for index in indices.split(","):
            overrides[int(index.strip())] = role
    return overrides


def mode(values: list) -> object | None:
    """@returns the most common value, ties broken by first appearance."""
    filtered = [value for value in values if value is not None]
    if not filtered:
        return None
    return Counter(filtered).most_common(1)[0][0]


def extract(pptx_path: Path, overrides: dict[int, str]) -> dict:
    """@returns the numeric profile of the deck at `pptx_path`."""
    presentation = Presentation(str(pptx_path))
    width = int(presentation.slide_width)
    height = int(presentation.slide_height)
    slides = list(presentation.slides)
    runs_per_slide = [runs_with_style(slide_runs(slide)) for slide in slides]
    roles = map_roles(runs_per_slide, len(slides), overrides)

    title_runs: dict[str, list[dict]] = defaultdict(list)
    body_runs: dict[str, list[dict]] = defaultdict(list)
    watermark_runs: dict[str, list[dict]] = defaultdict(list)
    slide_body: dict[int, list[dict]] = {}
    for offset, (role, runs) in enumerate(zip(roles, runs_per_slide)):
        for run in runs:
            light = luminance(run["color"] or "#000000") > 0.85
            if light and (run["sizePt"] or 0) >= 60:
                watermark_runs[role].append(run)
        dark = [run for run in runs if luminance(run["color"] or "#000000") <= 0.75]
        if dark:
            title = max(dark, key=lambda run: (run["sizePt"] or 0, -run["y"]))
            title_runs[role].append(title)
            remaining = [run for run in dark if run is not title]
            body_runs[role].extend(remaining)
            slide_body[offset] = remaining

    type_scale: dict[str, dict] = {}
    for role in ROLES:
        entry: dict[str, dict] = {}
        titles = title_runs.get(role, [])
        if titles:
            title = max(titles, key=lambda run: (run["sizePt"] or 0, -run["y"]))
            entry["title"] = {
                "sizePt": round(title["sizePt"] or 0, 1),
                "bold": bool(title["bold"]),
                "color": title["color"],
                "font": title["font"],
            }
        bodies = body_runs.get(role, [])
        if bodies:
            body = mode([(run["sizePt"], run["color"], run["font"]) for run in bodies])
            if body is not None:
                entry["body"] = {"sizePt": body[0], "color": body[1], "font": body[2]}
        if entry:
            type_scale[role] = entry

    # Role geometry: the title anchor, the number of text columns, and the gap
    # between them. Columns are clustered by shape x with a 1 in tolerance, and
    # the gap is measured from each cluster's right edge to the next cluster's
    # left edge using the authored shape widths.
    geometry: dict[str, dict] = {}
    for role in ROLES:
        titles = title_runs.get(role, [])
        if not titles:
            continue
        title = max(titles, key=lambda run: (run["sizePt"] or 0, -run["y"]))
        entry = {
            "titlePos": {"x": round(title["x"] / EMU_PER_INCH, 2), "y": round(title["y"] / EMU_PER_INCH, 2)},
        }
        if role in ("toc", "content"):
            # Per slide, cluster the body runs by x; the role keeps the most common
            # column count and the smallest positive gap between neighbouring
            # clusters (a negative figure means the authored text boxes overlap, so
            # it is not a card gap).
            counts: list[int] = []
            gaps: list[float] = []
            for offset, slide_role in enumerate(roles):
                if slide_role != role:
                    continue
                clusters: list[list[dict]] = []
                for run in sorted(slide_body.get(offset, []), key=lambda candidate: candidate["x"]):
                    if clusters and run["x"] - clusters[-1][-1]["x"] <= EMU_PER_INCH:
                        clusters[-1].append(run)
                    else:
                        clusters.append([run])
                counts.append(max(1, len(clusters)))
                for left, right in zip(clusters, clusters[1:]):
                    right_edge = max(run["x"] + run["w"] for run in left)
                    gaps.append((min(run["x"] for run in right) - right_edge) / EMU_PER_INCH)
            entry["columns"] = mode(counts) or 1
            entry["columnsMax"] = max(counts) if counts else 1
            positive = [gap for gap in gaps if gap > 0.05]
            if positive:
                entry["cardGapIn"] = round(min(positive), 2)
        else:
            entry["columns"] = 1
        if watermark_runs.get(role):
            watermark = max(watermark_runs[role], key=lambda run: run["sizePt"] or 0)
            entry["watermarkSizePt"] = round(watermark["sizePt"] or 0, 1)
        geometry[role] = entry

    palette = {
        "bg": "#FFFFFF",
        "title": mode([title["color"] for titles in title_runs.values() for title in titles]) or "#000000",
        "accent": mode([run["color"] for run in (run for runs in body_runs.values() for run in runs) if not is_grey(run["color"])]) or "#000000",
        "body": mode([run["color"] for run in (run for runs in body_runs.values() for run in runs) if is_grey(run["color"]) and (run["sizePt"] or 0) <= 16]) or "#262626",
        "muted": mode([run["color"] for run in (run for runs in body_runs.values() for run in runs) if is_grey(run["color"]) and 16 < (run["sizePt"] or 0) <= 30]) or "#595959",
        "watermark": mode([run["color"] for runs in watermark_runs.values() for run in runs]) or "#F2F7FA",
        "onAccent": mode([run["color"] for run in (run for runs in body_runs.values() for run in runs) if run["color"] and luminance(run["color"]) > 0.9]) or "#FFFFFF",
    }

    fonts = {
        "heading": mode([title["font"] for titles in title_runs.values() for title in titles]) or "Arial",
        "body": mode([run["font"] for run in (run for runs in body_runs.values() for run in runs)]) or "Arial",
        "number": mode([run["font"] for runs in runs_per_slide for run in runs if is_number_like(run["text"]) and (run["sizePt"] or 0) >= 32]) or "Arial",
    }

    background_mode = "photo" if any(has_full_bleed_picture(slide, width, height) for slide in slides) else "flat"

    page_number = False
    meta_footer = False
    for role, runs in zip(roles, runs_per_slide):
        for run in runs:
            text = run["text"].strip()
            if is_number_like(text) and (run["sizePt"] or 99) <= 18 and run["y"] > height * 0.85:
                page_number = True
            if role in ("cover", "ending") and (run["sizePt"] or 99) <= 18 and run["y"] > height * 0.8:
                meta_footer = True

    return {
        "version": PROFILE_VERSION,
        "canvas": {"widthEmu": width, "heightEmu": height},
        "palette": palette,
        "fonts": fonts,
        "typeScale": {role: type_scale[role] for role in ROLES if role in type_scale},
        "roles": {role: geometry[role] for role in ROLES if role in geometry},
        "chrome": {
            "sectionMarker": bool(watermark_runs),
            "metaFooter": meta_footer,
            "pageNumber": page_number,
        },
        "background": {"mode": background_mode, "overlayOpacity": 0.15 if background_mode == "photo" else 0.0},
    }


def copy_media(pptx_path: Path, destination: Path) -> int:
    """Copy the package's media into `destination`; @returns how many files were written."""
    destination.mkdir(parents=True, exist_ok=True)
    written = 0
    with zipfile.ZipFile(pptx_path) as archive:
        for name in archive.namelist():
            if not name.startswith("ppt/media/") or name.endswith("/"):
                continue
            target = destination / Path(name).name
            with archive.open(name) as source, target.open("wb") as sink:
                shutil.copyfileobj(source, sink)
            written += 1
    return written


def main(argv: list[str]) -> int:
    """@returns the process exit code."""
    parser = argparse.ArgumentParser(description="Extract a numeric design profile from a .pptx")
    parser.add_argument("--input", required=True, help="source .pptx (read-only)")
    parser.add_argument("--output", required=True, help="profile JSON to write")
    parser.add_argument("--roles", default=None, help="role overrides: role=1,2;role=3 or a JSON index map")
    parser.add_argument("--copy-media", default=None, help="explicitly copy ppt/media into this directory (default off)")
    options = parser.parse_args(argv)

    source = Path(options.input)
    if not source.is_file():
        print(f"design-profile: input is not a file: {options.input}", file=sys.stderr)
        return 2
    profile = extract(source, parse_roles(options.roles))
    output = Path(options.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(profile, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    copied = copy_media(source, Path(options.copy_media)) if options.copy_media else 0
    print(json.dumps({"output": str(output), "mediaCopied": copied}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
